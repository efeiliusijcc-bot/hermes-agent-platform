from io import BytesIO

from docx import Document
from openpyxl import Workbook

from hermes_knowledge_service.parsers import parse_document


def build_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(value)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii")
    )
    return bytes(output)


assert "MARKDOWN_FORMAT_SIGNAL" in parse_document("test.md", b"# MARKDOWN_FORMAT_SIGNAL").text

docx_buffer = BytesIO()
document = Document()
document.add_paragraph("WORD_FORMAT_SIGNAL")
document.save(docx_buffer)
assert "WORD_FORMAT_SIGNAL" in parse_document("test.docx", docx_buffer.getvalue()).text

xlsx_buffer = BytesIO()
workbook = Workbook()
workbook.active["A1"] = "EXCEL_FORMAT_SIGNAL"
workbook.save(xlsx_buffer)
workbook.close()
assert "EXCEL_FORMAT_SIGNAL" in parse_document("test.xlsx", xlsx_buffer.getvalue()).text

assert "PDF_FORMAT_SIGNAL" in parse_document("test.pdf", build_pdf("PDF_FORMAT_SIGNAL")).text
print("Phase 7 parser format validation passed")
